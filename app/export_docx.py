"""تصدير العرض إلى ملف Word احترافي بهوية عزوم — عربي RTL كامل."""
from contextvars import ContextVar

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

from .config import BRAND
from .proposal_builder import client_facing_pricing, flatten_boq_rows

# هوية المستند الحالي في ContextVar — كل طلب/خيط يرى هويته وحدها. كانت متغيرات
# وحدة عامة تتبادلها الطلبات المتزامنة فيخرج ملف بلون شركة أخرى.
_BRAND_CTX: ContextVar[tuple] = ContextVar("docx_brand", default=None)


def _resolve_brand(settings: dict) -> tuple:
    """(لون أساسي، لون مساند) كنص HEX. عزوم: أخضر الشعار المساند؛
    المستأجر: لونه هو للاثنين — لا يتسرب أخضر عزوم إلى وثائق غيرها."""
    hexv = (settings.get("_brand_color") or BRAND["primary"]).lstrip("#").upper()
    try:
        RGBColor.from_string(hexv)
    except Exception:
        hexv = BRAND["primary"]
    accent = BRAND["accent"] if _is_azoom(settings) else hexv
    return hexv, accent


def _primary_hex() -> str:
    b = _BRAND_CTX.get()
    return b[0] if b else BRAND["primary"]


def _accent_hex() -> str:
    b = _BRAND_CTX.get()
    return b[1] if b else BRAND["accent"]


def _primary() -> RGBColor:
    return RGBColor.from_string(_primary_hex())


def _accent() -> RGBColor:
    return RGBColor.from_string(_accent_hex())


def _is_azoom(settings: dict) -> bool:
    """هوية عزوم الثابتة (تذييل/اسم إنجليزي) لمستأجر عزوم نفسه (الشركة 1) فقط —
    لا بمطابقة السجل التجاري: أي مستأجر يكتب سجل عزوم لا يرث تذييلها."""
    return settings.get("_company_id") == 1


FONT = "Sakkal Majalla"
FONT_FALLBACK = "Arial"


def _set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)


def _run(paragraph, text, size=13, bold=False, color=None):
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT_FALLBACK
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:cs"), FONT)
    rFonts.set(qn("w:ascii"), FONT_FALLBACK)
    cs_bold = OxmlElement("w:bCs")
    cs_bold.set(qn("w:val"), "1" if bold else "0")
    rPr.append(cs_bold)
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), str(size * 2))
    rPr.append(sz_cs)
    if color:
        run.font.color.rgb = color
    return run


def _para(doc, text="", size=13, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    _set_rtl(p)
    if text:
        _run(p, text, size=size, bold=bold, color=color)
    return p


def _heading(doc, text, level=1):
    if level == 1:
        p = _para(doc, text, size=17, bold=True, color=_primary(), space_after=10)
        _add_bottom_border(p)
    else:
        _para(doc, text, size=14, bold=True, color=_accent(), space_after=8)


def _add_bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:color"), _accent_hex())
    borders.append(bottom)
    pPr.append(borders)


def _shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _table(doc, headers, rows, widths=None, money_cols=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # اتجاه الجدول من اليمين لليسار
    tblPr = table._tbl.tblPr
    bidi = OxmlElement("w:bidiVisual")
    tblPr.append(bidi)

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_rtl(p)
        _run(p, h, size=11, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        _shade_cell(cell, _primary_hex())

    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_rtl(p)
            txt = f"{value:,.2f}" if i in money_cols and isinstance(value, (int, float)) else str(value)
            _run(p, txt, size=11)
    if widths:
        for i, w in enumerate(widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    return table


def _page_footer(doc, settings=None):
    """تذييل رسمي في كل الصفحات — كما في عروض عزوم المطبوعة."""
    settings = settings or {}
    if _is_azoom(settings) or not settings:
        footer_text = f'{BRAND["footer_phone"]}   |   {BRAND["footer_address"]}   |   C.R. {BRAND["footer_cr"]}'
        footer_name = BRAND["name_en"]
    else:
        parts = [settings.get("company_phone", ""), settings.get("company_address", "")]
        if settings.get("company_cr"):
            parts.append(f'C.R. {settings["company_cr"]}')
        footer_text = "   |   ".join(x for x in parts if x) or settings.get("company_name", "")
        footer_name = settings.get("company_name", "")
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), "8")
        top.set(qn("w:color"), _accent_hex())
        borders.append(top)
        pPr.append(borders)
        _run(p, footer_name, size=9, bold=True, color=_primary())
        _run(p, "   —   ", size=9, color=_accent())
        _run(p, footer_text, size=9, color=RGBColor(0x66, 0x66, 0x66))


def _cover_page(doc, proposal, settings):
    for _ in range(2):
        doc.add_paragraph()
    # شعار الشركة في صدر الغلاف — هوية كل مستأجر (WEBP غير مدعوم في Word فيُتخطى)
    logo = settings.get("_logo_path", "")
    if logo and not logo.lower().endswith(".webp"):
        try:
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.add_run().add_picture(logo, width=Cm(4.2))
        except Exception:
            pass
    doc.add_paragraph()
    company_name = settings.get("company_name") or (BRAND["name_ar"] if _is_azoom(settings) else "")
    _para(doc, company_name, size=32, bold=True, color=_primary(), align=WD_ALIGN_PARAGRAPH.CENTER)
    if _is_azoom(settings):
        _para(doc, BRAND["name_en"], size=16, bold=True, color=_accent(), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30)
    else:
        _para(doc, "", size=10, space_after=24)
    _para(doc, "العرض الفني والمالي", size=26, bold=True, color=_primary(), align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, proposal["title"], size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=20)
    _para(doc, f"مقدم إلى: {proposal['client']}", size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, f"رقم العرض: {proposal['ref_no']}", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, f"تاريخ التقديم: {proposal['created_at'][:10]}", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, f"سريان العرض: {settings.get('validity_days', '90')} يوماً", size=12,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=40)
    _para(doc, settings.get("company_name", ""), size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    contact = " — ".join(x for x in (settings.get("company_phone"), settings.get("company_email")) if x)
    if contact:
        _para(doc, contact, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def export_proposal_docx(proposal: dict, settings: dict, path):
    """يبني ملف Word ويحفظه في `path` (مسار أو كائن ملف). آمن للتزامن."""
    token = _BRAND_CTX.set(_resolve_brand(settings))
    try:
        return _build_docx(proposal, settings, path)
    finally:
        _BRAND_CTX.reset(token)


def _build_docx(proposal: dict, settings: dict, path):
    data = proposal["data"]
    doc = Document()

    # هوامش وإعداد عام
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.right_margin = Cm(2.2)
        section.left_margin = Cm(2.2)

    # خصائص الملف باسم الشركة — لا يظهر اسم أي مكتبة برمجية في Author/Company
    from .quality_agent import clean_file_properties
    props = clean_file_properties(settings)
    doc.core_properties.author = props["author"]
    doc.core_properties.last_modified_by = props["last_modified_by"]
    doc.core_properties.title = proposal.get("title", "")
    doc.core_properties.subject = "عرض فني ومالي"
    doc.core_properties.comments = ""
    _page_footer(doc, settings)
    _cover_page(doc, proposal, settings)

    # ---------- الجزء الأول: العرض الفني ----------
    _heading(doc, "أولاً: العرض الفني")
    for sec in data.get("technical_sections", []):
        _heading(doc, sec["title"], level=2)
        for paragraph_text in sec["body"].split("\n"):
            if paragraph_text.strip():
                _para(doc, paragraph_text.strip(), size=12)

    scope = data.get("scope") or []
    if scope and len(scope) > 1:
        _heading(doc, "نطاق العمل التفصيلي", level=2)
        for item in scope:
            _para(doc, f"•  {item}", size=12)

    team = data.get("team") or []
    if team:
        _heading(doc, "فريق العمل المقترح", level=2)
        _table(doc, ["الدور الوظيفي", "العدد"],
               [(t["role"], int(t["count"])) for t in team], widths=[10, 4])

    matrix = data.get("compliance_matrix") or []
    if matrix:
        _heading(doc, "مصفوفة الالتزام بالمتطلبات", level=2)
        _table(doc, ["المتطلب", "الالتزام", "الموضع في العرض"],
               [(m["requirement"], m["response"], m["reference"]) for m in matrix],
               widths=[8, 3, 5])

    doc.add_page_break()

    # ---------- الجزء الثاني: الخطة التنفيذية ----------
    _heading(doc, "ثانياً: الخطة التنفيذية للمشروع")
    _para(doc, f"المدة الإجمالية المقترحة: {int(data.get('duration_weeks', 0))} أسبوعاً", size=13, bold=True)
    for i, phase in enumerate(data.get("plan", []), 1):
        _heading(doc, f"المرحلة {i}: {phase['phase']} ({int(phase['duration_weeks'])} أسابيع)", level=2)
        _para(doc, phase["description"], size=12)
        for d in phase.get("deliverables", []):
            _para(doc, f"–  {d}", size=11)

    doc.add_page_break()

    # ---------- الجزء الثالث: العرض المالي ----------
    # أسعار البنود محمَّلة بكامل التكاليف — لا تظهر أي نسب داخلية للعميل
    _heading(doc, "ثالثاً: العرض المالي")
    _heading(doc, "جدول الكميات والأسعار", level=2)
    boq = client_facing_pricing(data.get("boq", []), data.get("financial", {}))
    _table(
        doc,
        ["م", "الكود", "البند", "الوحدة", "الكمية", "سعر الوحدة (ر.س)", "الإجمالي (ر.س)"],
        flatten_boq_rows(boq),
        widths=[1, 2, 6, 2, 2, 3, 3],
        money_cols=(5, 6),
    )

    fin = data.get("financial", {})
    _heading(doc, "ملخص القيمة الإجمالية", level=2)
    _table(
        doc,
        ["البيان", "القيمة (ريال سعودي)"],
        [
            ("إجمالي جدول الكميات والأسعار", fin.get("subtotal", 0)),
            (f"ضريبة القيمة المضافة ({fin.get('vat_rate', 15):g}%)", fin.get("vat", 0)),
            ("الإجمالي النهائي شامل الضريبة", fin.get("grand_total", 0)),
        ],
        widths=[9, 5],
        money_cols=(1,),
    )

    if (settings.get("payment_terms") or "").strip():   # لا عنوان فارغ لمستأجر لم يحدد شروطه
        _heading(doc, "شروط الدفع", level=2)
        _para(doc, settings["payment_terms"], size=12)

    assumptions = data.get("assumptions") or []
    if assumptions:
        _heading(doc, "الافتراضات والاستثناءات", level=2)
        for a in assumptions:
            _para(doc, f"•  {a}", size=11)

    _para(doc)
    _para(doc, "وتفضلوا بقبول فائق الاحترام والتقدير،", size=12)
    _para(doc, settings.get("company_name", ""), size=13, bold=True, color=_primary())

    doc.save(path)
    return path
